import streamlit as st
import fitz
from docx import Document
from deep_translator import GoogleTranslator
from langdetect import detect
import tempfile
import os
from pathlib import Path
import time


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="DocMorph",
    page_icon="🌐",
    layout="wide"
)


# ============================================================
# CUSTOM CSS
# ============================================================

st.markdown(
    """
    <style>

    .main-title {
        font-size: 42px;
        font-weight: 700;
        text-align: center;
        margin-bottom: 5px;
    }

    .subtitle {
        text-align: center;
        font-size: 18px;
        margin-bottom: 30px;
    }

    .step-card {
        padding: 15px;
        border-radius: 12px;
        text-align: center;
        border: 1px solid #444;
    }

    </style>
    """,
    unsafe_allow_html=True
)


# ============================================================
# TITLE
# ============================================================

st.markdown(
    '<div class="main-title">🌐 DocMorph</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="subtitle">'
    'Smart Document Language Converter'
    '</div>',
    unsafe_allow_html=True
)


# ============================================================
# SESSION STATE
# ============================================================

if "extracted_pages" not in st.session_state:
    st.session_state.extracted_pages = []

if "source_language" not in st.session_state:
    st.session_state.source_language = ""

if "translated_file" not in st.session_state:
    st.session_state.translated_file = None

if "page_count" not in st.session_state:
    st.session_state.page_count = 0


# ============================================================
# WORKFLOW HEADER
# ============================================================

step1, step2, step3, step4 = st.columns(4)

with step1:
    st.info("01\n\n📥 INPUT")

with step2:
    st.info("02\n\n🔍 ANALYZE")

with step3:
    st.info("03\n\n🌐 TRANSLATE")

with step4:
    st.info("04\n\n📤 EXPORT")


# ============================================================
# UPLOAD SECTION
# ============================================================

st.subheader("📥 Document Input")

uploaded_file = st.file_uploader(
    "Upload your document",
    type=["pdf", "docx", "txt"],
    help="Supported formats: PDF, DOCX and TXT"
)


# ============================================================
# TARGET LANGUAGES
# ============================================================

# ONLY FOUR LANGUAGES FOR NOW

language_map = {
    "English": "en",
    "Kannada": "kn",
    "Telugu": "te",
    "Tamil": "ta"
}


# ============================================================
# DOCUMENT EXTRACTION
# ============================================================

def extract_pdf_pages(file_bytes):

    pages = []

    pdf = fitz.open(
        stream=file_bytes,
        filetype="pdf"
    )

    for page_number, page in enumerate(pdf):

        text = page.get_text("text")

        pages.append({
            "page": page_number + 1,
            "text": text.strip()
        })

    pdf.close()

    return pages


def extract_docx(file_bytes):

    temp_path = os.path.join(
        tempfile.gettempdir(),
        "input_document.docx"
    )

    with open(temp_path, "wb") as f:
        f.write(file_bytes)

    document = Document(temp_path)

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:

            paragraphs.append({
                "page": 1,
                "text": text
            })

    os.remove(temp_path)

    return paragraphs


def extract_txt(file_bytes):

    text = file_bytes.decode(
        "utf-8",
        errors="ignore"
    )

    return [{
        "page": 1,
        "text": text
    }]


# ============================================================
# CREATE TEXT CHUNKS
# ============================================================

def create_chunks(text, max_chars=2500):

    text = text.strip()

    if not text:
        return []

    chunks = []

    start = 0

    while start < len(text):

        end = min(
            start + max_chars,
            len(text)
        )

        # Try to break at a space
        if end < len(text):

            space_position = text.rfind(
                " ",
                start,
                end
            )

            if space_position > start:

                end = space_position

        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        start = end

    return chunks


# ============================================================
# TRANSLATE ONE BATCH
# ============================================================

def translate_batch_with_retry(
    texts,
    target_language,
    max_retries=3
):

    if not texts:
        return []

    for attempt in range(max_retries):

        try:

            translator = GoogleTranslator(
                source="auto",
                target=target_language
            )

            translated = translator.translate_batch(
                texts
            )

            return translated

        except Exception as e:

            if attempt < max_retries - 1:

                wait_time = (
                    2 ** attempt
                )

                time.sleep(wait_time)

            else:

                raise e


# ============================================================
# TRANSLATE DOCUMENT
# ============================================================

def translate_document_pages(
    pages,
    target_language
):

    # --------------------------------------------------------
    # Prepare all text
    # --------------------------------------------------------

    all_chunks = []
    chunk_page_numbers = []

    for item in pages:

        text = item["text"].strip()

        if not text:
            continue

        chunks = create_chunks(
            text,
            max_chars=2500
        )

        for chunk in chunks:

            all_chunks.append(chunk)

            chunk_page_numbers.append(
                item["page"]
            )

    if not all_chunks:

        return []

    # --------------------------------------------------------
    # Translate in batches
    # --------------------------------------------------------

    translated_chunks = []

    batch_size = 5

    total_batches = (
        len(all_chunks) + batch_size - 1
    ) // batch_size

    for start in range(
        0,
        len(all_chunks),
        batch_size
    ):

        batch = all_chunks[
            start:start + batch_size
        ]

        batch_number = (
            start // batch_size
        ) + 1

        # Show progress
        st.write(
            f"🌐 Translating batch "
            f"{batch_number} of "
            f"{total_batches}..."
        )

        translated_batch = (
            translate_batch_with_retry(
                batch,
                target_language
            )
        )

        translated_chunks.extend(
            translated_batch
        )

        # IMPORTANT:
        # Wait between batches to reduce
        # Google rate-limit problems.
        if (
            start + batch_size
            < len(all_chunks)
        ):

            time.sleep(1.5)

    # --------------------------------------------------------
    # Rebuild translated pages
    # --------------------------------------------------------

    page_text = {}

    for index, translated_text in enumerate(
        translated_chunks
    ):

        page_number = (
            chunk_page_numbers[index]
        )

        if page_number not in page_text:

            page_text[page_number] = []

        page_text[page_number].append(
            translated_text
        )

    translated_pages = []

    for page in pages:

        page_number = page["page"]

        translated_text = "\n\n".join(
            page_text.get(
                page_number,
                []
            )
        )

        translated_pages.append({
            "page": page_number,
            "translated": translated_text
        })

    return translated_pages


# ============================================================
# CREATE DOCX
# ============================================================

def create_translated_docx(
    translated_pages,
    source_language,
    target_language,
    original_name
):

    output_path = os.path.join(
        tempfile.gettempdir(),
        "DocMorph_Translated.docx"
    )

    document = Document()

    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    document.add_heading(
        "DocMorph - Translated Document",
        level=0
    )

    document.add_paragraph(
        f"Original Document: {original_name}"
    )

    document.add_paragraph(
        f"Source Language: {source_language}"
    )

    document.add_paragraph(
        f"Target Language: {target_language}"
    )

    document.add_paragraph(
        "Generated using DocMorph"
    )

    document.add_page_break()

    # --------------------------------------------------------
    # TRANSLATED PAGES
    # --------------------------------------------------------

    for item in translated_pages:

        document.add_heading(
            f"Page {item['page']}",
            level=1
        )

        if item["translated"]:

            document.add_paragraph(
                item["translated"]
            )

        # Add page break except after last page
        if item != translated_pages[-1]:

            document.add_page_break()

    document.save(output_path)

    return output_path


# ============================================================
# ANALYZE BUTTON
# ============================================================

if uploaded_file is not None:

    st.success(
        f"📄 {uploaded_file.name} uploaded"
    )

    col_a, col_b = st.columns(2)

    with col_a:

        if st.button(
            "🔍 Analyze Document",
            use_container_width=True
        ):

            file_bytes = (
                uploaded_file.getvalue()
            )

            extension = Path(
                uploaded_file.name
            ).suffix.lower()

            try:

                # PDF
                if extension == ".pdf":

                    pages = extract_pdf_pages(
                        file_bytes
                    )

                # DOCX
                elif extension == ".docx":

                    pages = extract_docx(
                        file_bytes
                    )

                # TXT
                elif extension == ".txt":

                    pages = extract_txt(
                        file_bytes
                    )

                else:

                    st.error(
                        "Unsupported file format."
                    )

                    st.stop()

                # Remove completely empty pages
                pages = [
                    page
                    for page in pages
                    if page["text"].strip()
                ]

                st.session_state.extracted_pages = (
                    pages
                )

                st.session_state.page_count = (
                    len(pages)
                )

                # ------------------------------------------------
                # DETECT LANGUAGE
                # ------------------------------------------------

                combined_text = " ".join(
                    item["text"]
                    for item in pages
                    if item["text"]
                )

                if combined_text:

                    try:

                        detected = detect(
                            combined_text[:3000]
                        )

                        st.session_state.source_language = (
                            detected
                        )

                    except Exception:

                        st.session_state.source_language = (
                            "Unknown"
                        )

                st.success(
                    "✅ Document analyzed successfully!"
                )

            except Exception as e:

                st.error(
                    f"❌ Analysis error: {str(e)}"
                )


# ============================================================
# DOCUMENT INFORMATION
# ============================================================

if st.session_state.extracted_pages:

    st.subheader(
        "🔍 Document Intelligence"
    )

    info1, info2, info3 = st.columns(3)

    with info1:

        st.metric(
            "Pages / Sections",
            st.session_state.page_count
        )

    with info2:

        st.metric(
            "Detected Language",
            st.session_state.source_language
        )

    with info3:

        total_chars = sum(
            len(item["text"])
            for item
            in st.session_state.extracted_pages
        )

        st.metric(
            "Characters",
            total_chars
        )


# ============================================================
# TARGET LANGUAGE
# ============================================================

if st.session_state.extracted_pages:

    st.subheader(
        "🌐 Translation Settings"
    )

    target_language_name = st.selectbox(
        "Choose Target Language",
        list(language_map.keys())
    )

    target_language = language_map[
        target_language_name
    ]

    # ========================================================
    # PREVIEW
    # ========================================================

    with st.expander(
        "👁 Preview Extracted Content"
    ):

        for item in (
            st.session_state.extracted_pages
        ):

            st.markdown(
                f"**Page / Section "
                f"{item['page']}**"
            )

            preview = item["text"]

            if len(preview) > 500:

                preview = (
                    preview[:500]
                    + "..."
                )

            st.write(preview)

    # ========================================================
    # TRANSLATE BUTTON
    # ========================================================

    if st.button(
        "🌐 Translate Document",
        type="primary",
        use_container_width=True
    ):

        progress = st.progress(0)

        status = st.empty()

        try:

            status.write(
                "🔄 Preparing document for translation..."
            )

            progress.progress(10)

            # ------------------------------------------------
            # TRANSLATE ALL PAGES
            # ------------------------------------------------

            translated_pages = (
                translate_document_pages(
                    st.session_state.extracted_pages,
                    target_language
                )
            )

            progress.progress(80)

            # ------------------------------------------------
            # CREATE DOCX
            # ------------------------------------------------

            status.write(
                "📄 Creating translated DOCX..."
            )

            output_path = create_translated_docx(
                translated_pages,
                st.session_state.source_language,
                target_language_name,
                uploaded_file.name
            )

            st.session_state.translated_file = (
                output_path
            )

            progress.progress(100)

            status.write(
                "✅ Translation completed!"
            )

            st.success(
                "🎉 Your translated document is ready!"
            )

        except Exception as e:

            st.error(
                f"❌ Translation error: {str(e)}"
            )


# ============================================================
# DOWNLOAD
# ============================================================

if st.session_state.translated_file:

    st.subheader("📤 Export")

    with open(
        st.session_state.translated_file,
        "rb"
    ) as file:

        document_bytes = file.read()

    st.download_button(
        label="⬇️ Download Translated DOCX",
        data=document_bytes,
        file_name="DocMorph_Translated.docx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        use_container_width=True
    )
